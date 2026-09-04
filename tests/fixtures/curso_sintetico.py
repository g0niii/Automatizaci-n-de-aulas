# -*- coding: utf-8 -*-
"""Construye un curso de prueba sintético, equivalente a lo que envía asesoría.

Por qué existe
--------------
El material real de asesoría (`Aulas a generar/`) no se versiona: es pesado y
tiene contenido de cátedra. Sin él, la mitad de la suite se salteaba y el
generador —`imscc_builder.py`, el corazón del proyecto— nunca se ejecutaba en
CI. Este módulo arma en un directorio temporal un curso mínimo pero completo,
con un archivo por cada rama del clasificador de `folder_scanner`.

Por qué generado y no commiteado
--------------------------------
Un .docx es un ZIP: commitearlo deja un blob opaco que nadie puede revisar en
un diff ni ajustar sin abrir Word. Generarlo por código lo vuelve legible,
diffeable y extensible — y usa python-docx/openpyxl, que ya son dependencias.

Qué cubre
---------
  módulos, foros, actividades, guiones de video, programa, figuras de diseño,
  esquemas, foto del docente, y DOS descartes (carpeta 'Devoluciones' y archivo
  'Borrador…') que verifican que el material no usable quede afuera.
"""

from pathlib import Path

import openpyxl
from docx import Document
from PIL import Image

CODIGO = "TEST001"
NOMBRE_CURSO = "Curso Sintético de Prueba"
DOCENTE = "Dra. Ana Prueba"

# La planilla manda: define qué páginas y qué recursos pide cada módulo.
# Formato "nuevo" (9 columnas, Planilla de montaje V3).
_FILAS_PLANILLA = [
    ["DATOS DE LA ASIGNATURA", "", "", "", "", "", "", "", ""],
    ["Código de la asignatura", CODIGO, "", "", "", "", "", "", ""],
    ["Nombre", NOMBRE_CURSO, "", "", "", "", "", "", ""],
    ["Contenidista", DOCENTE, "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", "", ""],
    ["Link Drive:", "", "", "", "", "", "", "", ""],
    ["Página de inicio", "Programa", "Programa de la asignatura", "PDF",
     "Descarga", "", "Completo", "", ""],
    # --- Módulo 1 -------------------------------------------------------
    # Ojo con la numeración: las páginas de CONTENIDOS van numeradas (así las
    # clasifica _clasificar_item como PAGINA y las matchea contra las secciones
    # del DOCX), mientras que los ítems de ACTIVIDADES van SIN numerar — si se
    # numeran, caen en la rama PAGINA y se buscan como sección inexistente.
    ["Módulo 1", "INTRODUCCIÓN", "", "", "", "", "", "", ""],
    ["", "Texto introductorio", "Introducción al módulo", "Lectura", "",
     "", "Completo", "", ""],
    ["", "CONTENIDOS", "", "", "", "", "", "", ""],
    ["", "1.1. Conceptos fundamentales", "Desarrollo teórico", "Lectura", "",
     "", "Completo", "", ""],
    ["", "1.2. Casos de aplicación", "Segunda página del módulo", "Lectura",
     "", "", "Completo", "", ""],
    ["", "ACTIVIDADES", "", "", "", "", "", "", ""],
    ["", "Foro de debate", "Foro del módulo 1", "Foro", "Grupal", "",
     "Completo", "", ""],
    ["", "Trabajo práctico", "Actividad del módulo 1", "Tarea",
     "Individual", "", "Completo", "", ""],
    # --- Módulo 2 -------------------------------------------------------
    ["Módulo 2", "INTRODUCCIÓN", "", "", "", "", "", "", ""],
    ["", "Texto introductorio", "Introducción al módulo 2", "Lectura",
     "", "", "Completo", "", ""],
    ["", "CONTENIDOS", "", "", "", "", "", "", ""],
    ["", "2.1. Profundización", "Desarrollo teórico del módulo 2", "Lectura",
     "", "", "Completo", "", ""],
]


def _png(path: Path, color: tuple, tam: tuple = (48, 48)) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", tam, color).save(path)


def _docx_modulo(path: Path, numero: int, titulo_modulo: str,
                 secciones: list) -> None:
    """DOCX de desarrollo teórico con la forma de la plantilla institucional.

    Reproduce lo que `docx_probe.perfilar_docx` espera encontrar:
      - tabla de metadatos inicial (de ahí sale el nombre del módulo),
      - bloques Introducción / Objetivos / Conclusiones / Bibliografía,
      - secciones numeradas "N.N. Título" con estilo de encabezado, que son
        las que el segmentador convierte en páginas.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    tabla = doc.add_table(rows=0, cols=2)
    for clave, valor in (("Nombre de la carrera", "Carrera de Prueba"),
                         ("Asignatura", NOMBRE_CURSO),
                         ("Docente", DOCENTE),
                         ("Nombre del módulo", titulo_modulo)):
        fila = tabla.add_row().cells
        fila[0].text = clave
        fila[1].text = valor

    doc.add_heading(f"Módulo {numero}", level=1)

    doc.add_heading("Introducción", level=2)
    doc.add_paragraph(f"Este módulo aborda {titulo_modulo.lower()} y su lugar "
                      "dentro del recorrido de la asignatura.")
    doc.add_heading("Objetivos", level=2)
    doc.add_paragraph("Comprender los conceptos centrales del módulo.")
    doc.add_paragraph("Aplicarlos al análisis de situaciones concretas.")

    for titulo, parrafos in secciones:
        doc.add_heading(titulo, level=2)
        for p in parrafos:
            doc.add_paragraph(p)

    doc.add_heading("Conclusiones", level=2)
    doc.add_paragraph("Cierre de los contenidos trabajados en el módulo.")
    doc.add_heading("Bibliografía", level=2)
    doc.add_paragraph("Autor, A. (2020). Título del texto de referencia. Editorial.")
    doc.save(str(path))


def _docx_simple(path: Path, titulo: str, parrafos: list) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(titulo, level=1)
    for p in parrafos:
        doc.add_paragraph(p)
    doc.save(str(path))


def _planilla(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estructura"
    for fila in _FILAS_PLANILLA:
        ws.append(fila)
    wb.save(str(path))
    wb.close()


def construir(destino: Path) -> Path:
    """Arma el curso sintético dentro de `destino`. Devuelve la carpeta raíz
    del curso (la que se le pasa a `escanear`)."""
    raiz = Path(destino) / NOMBRE_CURSO
    raiz.mkdir(parents=True, exist_ok=True)

    # 1 - Programa
    _docx_simple(raiz / "1- Programa" / "Programa de la asignatura.docx",
                 "Programa de la asignatura",
                 ["Fundamentación de la asignatura de prueba.",
                  "Objetivos generales y contenidos mínimos."])

    # 2 - Desarrollo teórico por módulos
    desarrollo = raiz / "2- Desarrollo teórico por módulos"
    _docx_modulo(
        desarrollo / "Material multimedial modular - Módulo 1.docx", 1,
        "Fundamentos de la materia",
        [("1.1. Conceptos fundamentales",
          ["Definición de los conceptos centrales del módulo.",
           "Cada concepto se ilustra con ejemplos de la práctica profesional."]),
         ("1.2. Casos de aplicación",
          ["Análisis de situaciones reales donde se aplican los conceptos.",
           "Se propone contrastar los casos con la bibliografía."])])
    _docx_modulo(
        desarrollo / "Material multimedial modular - Módulo 2.docx", 2,
        "Profundización y cierre",
        [("2.1. Profundización",
          ["Desarrollo ampliado de los conceptos y su discusión actual.",
           "Cierre integrador del recorrido propuesto."]),
         ("2.2. Síntesis del recorrido",
          ["Recapitulación de los ejes trabajados en los dos módulos."])])

    # 3 - Actividades y foros
    actividades = raiz / "3- Actividades y foros"
    # Los nombres replican el título que la planilla le da a cada recurso
    # ("Foro del módulo 1", "Actividad del módulo 1"). No es cosmético: el
    # generador se niega a cargar foros y tareas cuyo match con el DOCX quede
    # por debajo del 60% de confianza (imscc_builder._inyectar_foros_y_actividades),
    # así que un nombre flojo deja el recurso vacío y el test no probaría nada.
    _docx_simple(actividades / "Foro del módulo 1 - Debate inicial.docx",
                 "Foro de debate",
                 ["Consigna: a partir de la lectura del módulo, comparta una "
                  "reflexión sobre los conceptos trabajados.",
                  "Responda al menos a un aporte de otro compañero."])
    _docx_simple(actividades / "Actividad del módulo 1 - Trabajo práctico.docx",
                 "Trabajo práctico 1",
                 ["Consigna: elabore un informe breve aplicando los conceptos "
                  "del módulo a un caso de su elección.",
                  "Extensión sugerida: dos páginas."])
    _docx_simple(actividades / "Guion video Módulo 1.docx",
                 "Guion del video introductorio",
                 ["Presentación del módulo a cargo de la docente."])

    # 4 - Diseño (figuras que reemplazan a las imágenes del docente)
    diseno = raiz / "4- Diseño"
    _png(diseno / "Figura 1 - Mapa conceptual.png", (120, 30, 60))
    _png(diseno / "Esquema de decisión.png", (30, 90, 140))

    # 6 - Maquetación (planilla + foto del docente)
    maquetacion = raiz / "6- Maquetación"
    _planilla(maquetacion / "Estructura general - Para maquetación.xlsx")
    _png(maquetacion / "Foto de la docente.jpg", (200, 190, 170), (64, 64))

    # --- Material que NO debe entrar --------------------------------------
    # Carpeta descartada por nombre.
    _docx_simple(raiz / "Devoluciones" / "Módulo 1 revisado.docx",
                 "Versión con devoluciones", ["No debe usarse como fuente."])
    # Archivo descartado por nombre, aunque esté en la carpeta correcta.
    _docx_simple(desarrollo / "Borrador - Módulo 3.docx",
                 "Borrador descartado", ["No debe usarse como fuente."])

    return raiz
