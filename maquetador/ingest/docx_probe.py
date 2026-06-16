# -*- coding: utf-8 -*-
"""Detección y segmentación de los DOCX de desarrollo teórico.

Los docentes entregan los documentos con criterios dispares: algunos usan
estilos de Word (Heading 1/2, Title), otros solo negritas, otros mezclan.
En vez de asumir un formato, se prueban varias ESTRATEGIAS de detección de
títulos de sección y se elige la que produce el resultado más consistente.

Estrategias (en orden de confiabilidad):
  1. heading_styles : párrafos con estilo Heading/Título/Title numerados
  2. bold_numbered  : párrafos en negrita que arrancan con numeración N.N
  3. plain_numbered : cualquier párrafo corto que arranca con numeración N.N

Además se recolectan CANDIDATOS sin numerar (párrafos cortos en negrita o
con estilo de título): hay docentes que escriben los títulos de sección sin
número, y solo se pueden asociar a la planilla por similitud de texto.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

_PAT_SECCION = re.compile(r"^(\d+(?:\.\d+)+)\.?\s+(.{2,})")
_MAX_LARGO_TITULO = 150


@dataclass
class SeccionDetectada:
    numero: str          # "1.1", "1.4.2"…
    titulo: str          # texto sin la numeración
    texto_completo: str
    indice_parrafo: int

    def to_dict(self):
        return {"numero": self.numero, "titulo": self.titulo}


@dataclass
class CandidatoSinNumero:
    titulo: str
    indice_parrafo: int

    def to_dict(self):
        return {"titulo": self.titulo}


@dataclass
class PerfilDocx:
    archivo: Path
    estrategia: str = ""           # la ganadora
    secciones: list = field(default_factory=list)   # list[SeccionDetectada]
    candidatos: list = field(default_factory=list)  # list[CandidatoSinNumero]
    metadatos: dict = field(default_factory=dict)   # tabla inicial del DOCX
    tiene_intro: bool = False
    tiene_objetivos: bool = False
    tiene_conclusion: bool = False
    tiene_referencias: bool = False
    detalle_estrategias: dict = field(default_factory=dict)  # {nombre: n_secciones}

    def to_dict(self):
        return {"archivo": self.archivo.name, "estrategia": self.estrategia,
                "secciones": [s.to_dict() for s in self.secciones],
                "candidatos_sin_numero": [c.to_dict() for c in self.candidatos],
                "intro": self.tiene_intro, "objetivos": self.tiene_objetivos,
                "conclusion": self.tiene_conclusion,
                "referencias": self.tiene_referencias,
                "detalle_estrategias": self.detalle_estrategias}


def _es_estilo_titulo(parrafo) -> bool:
    estilo = parrafo.style.name if parrafo.style else ""
    return any(s in estilo for s in ("Heading", "Título", "Title"))


def _es_negrita(parrafo) -> bool:
    runs = [r for r in parrafo.runs if r.text.strip()]
    if not runs:
        return False
    en_negrita = sum(1 for r in runs if r.bold)
    return en_negrita >= len(runs) / 2


def _detectar(parrafos, criterio) -> list:
    """Aplica un criterio (función parrafo->bool) y devuelve las secciones
    numeradas que cumplen el patrón N.N + criterio."""
    secciones = []
    for i, p in enumerate(parrafos):
        texto = p.text.strip()
        if not texto or len(texto) > _MAX_LARGO_TITULO:
            continue
        m = _PAT_SECCION.match(texto)
        if m and criterio(p):
            secciones.append(SeccionDetectada(
                numero=m.group(1), titulo=m.group(2).strip(),
                texto_completo=texto, indice_parrafo=i))
    return secciones


def perfilar_docx(path: Path) -> PerfilDocx:
    """Analiza un DOCX de módulo y devuelve su perfil: estrategia de
    segmentación aplicable y secciones detectadas."""
    doc = Document(str(path))
    parrafos = doc.paragraphs
    perfil = PerfilDocx(archivo=Path(path))

    # Tabla de metadatos con que arranca la plantilla institucional:
    # "Nombre de la carrera / Asignatura / Docente / Nombre del módulo"
    if doc.tables:
        for fila in doc.tables[0].rows:
            celdas = [c.text.strip() for c in fila.cells]
            if len(celdas) >= 2 and celdas[0] and celdas[1]:
                clave = re.sub(r"\s+", " ", celdas[0].lower())
                perfil.metadatos[clave] = re.sub(r"\s+", " ", celdas[1]).strip()

    estrategias = [
        ("heading_styles", lambda p: _es_estilo_titulo(p)),
        ("bold_numbered", lambda p: _es_negrita(p)),
        ("plain_numbered", lambda p: True),
    ]

    resultados = {}
    for nombre, criterio in estrategias:
        resultados[nombre] = _detectar(parrafos, criterio)
        perfil.detalle_estrategias[nombre] = len(resultados[nombre])

    # Elegir: la primera estrategia (más confiable) que detecte al menos 2
    # secciones; si ninguna llega a 2, la que más detecte.
    for nombre, _ in estrategias:
        if len(resultados[nombre]) >= 2:
            perfil.estrategia = nombre
            perfil.secciones = resultados[nombre]
            break
    else:
        mejor = max(resultados, key=lambda k: len(resultados[k]))
        perfil.estrategia = mejor if resultados[mejor] else "ninguna"
        perfil.secciones = resultados[mejor]

    # Marcadores especiales y candidatos sin numerar
    _NO_CANDIDATO = re.compile(
        r"^(figura|tabla|esquema|contenido|destaquemos|clic |para reflexionar|"
        r"te invito|importante|record[áa]|atenci[óo]n|nota:)", re.I)
    indices_numerados = {s.indice_parrafo for s in perfil.secciones}

    for i, p in enumerate(parrafos):
        t = p.text.strip().lower()
        texto = p.text.strip()
        if not texto or len(texto) > _MAX_LARGO_TITULO:
            continue
        es_marcador = _es_estilo_titulo(p) or _es_negrita(p)
        if not es_marcador:
            continue
        if t.startswith("introducci"):
            perfil.tiene_intro = True
        elif t.startswith("objetivo") or " objetivos" in t[:30]:
            perfil.tiene_objetivos = True
        elif t.startswith("conclusi") or "cierre" in t or "reflexión final" in t:
            perfil.tiene_conclusion = True
        elif t.startswith("referencia") or t.startswith("bibliograf"):
            perfil.tiene_referencias = True
        elif i not in indices_numerados and not _NO_CANDIDATO.match(texto) \
                and len(texto) >= 15:
            perfil.candidatos.append(CandidatoSinNumero(titulo=texto,
                                                        indice_parrafo=i))

    return perfil
