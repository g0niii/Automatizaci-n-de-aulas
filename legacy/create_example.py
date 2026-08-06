"""
Script de ejemplo — Genera archivos de prueba en input/ para validar el generador IMSCC.
Crea un archivo XLSX de estructura y documentos DOCX de ejemplo con contenido real.

Uso:
    python create_example.py
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from config import INPUT_DIR, MEDIA_DIR


def create_example_xlsx(filepath: Path):
    """Crea el archivo XLSX de ejemplo con estructura de curso."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estructura"

    # Encabezados
    headers = ["modulo", "orden", "tipo", "titulo", "archivo_docx", "archivo_media", "plantilla", "url", "puntos", "intentos", "notas"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = openpyxl.styles.Font(bold=True)

    # Datos de ejemplo
    data = [
        # Módulo 0: Presentación
        ["Presentación del curso", 1, "pagina", "Bienvenida al curso", "bienvenida.docx", "", "home_page", "", 0, 0, "Página de inicio"],
        ["Presentación del curso", 2, "foro", "Foro de presentación", "", "", "", "", 0, 0, "Preséntate aquí"],
        ["Presentación del curso", 3, "pagina", "Programa de la asignatura", "programa.docx", "", "cidilabs_page", "", 0, 0, ""],

        # Módulo 1: Fundamentos
        ["Módulo 1: Fundamentos", 1, "pagina", "Introducción M1", "tema1_intro.docx", "", "cidilabs_page", "", 0, 0, ""],
        ["Módulo 1: Fundamentos", 2, "pagina", "Desarrollo teórico", "tema1_desarrollo.docx", "", "cidilabs_page", "", 0, 0, ""],
        ["Módulo 1: Fundamentos", 3, "tarea", "Actividad obligatoria M1", "actividad1.docx", "", "cidilabs_page", "", 10, 2, "Consigna de tarea"],
        ["Módulo 1: Fundamentos", 4, "evaluacion", "Autoevaluación M1", "", "", "", "", 10, 3, "Quiz de opción múltiple"],
        ["Módulo 1: Fundamentos", 5, "url_externa", "Video: Introducción", "", "", "", "https://www.youtube.com/watch?v=ejemplo", 0, 0, ""],

        # Módulo 2: Profundización
        ["Módulo 2: Profundización", 1, "pagina", "Análisis avanzado", "tema2_analisis.docx", "", "cidilabs_page", "", 0, 0, ""],
        ["Módulo 2: Profundización", 2, "foro", "Foro de debate M2", "tema2_casos.docx", "", "cidilabs_page", "", 0, 0, "Debate grupal"],
        ["Módulo 2: Profundización", 3, "tarea", "Trabajo integrador final", "actividad2.docx", "", "cidilabs_page", "", 100, 1, "Entrega final"],
    ]

    for row_idx, row_data in enumerate(data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Ajustar anchos de columna
    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 40)

    wb.save(str(filepath))
    print(f"  [OK] Creado: {filepath}")


def create_example_docx(filepath: Path, title: str, content_paragraphs: list[str]):
    """Crea un archivo DOCX de ejemplo."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    doc = Document()

    # Título
    heading = doc.add_heading(title, level=1)

    # Párrafos de contenido
    for para_text in content_paragraphs:
        if para_text.startswith("## "):
            doc.add_heading(para_text[3:], level=2)
        elif para_text.startswith("### "):
            doc.add_heading(para_text[4:], level=3)
        elif para_text.startswith("- "):
            doc.add_paragraph(para_text[2:], style="List Bullet")
        elif para_text.startswith("1. "):
            doc.add_paragraph(para_text[3:], style="List Number")
        else:
            p = doc.add_paragraph(para_text)

    doc.save(str(filepath))
    print(f"  [OK] Creado: {filepath}")


def main():
    print()
    print("=" * 50)
    print("  Generador de archivos de ejemplo")
    print("=" * 50)
    print()

    # Crear directorios
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    MEDIA_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Crear XLSX de estructura
    create_example_xlsx(INPUT_DIR / "estructura.xlsx")

    # 2. Crear documentos DOCX de ejemplo
    docx_files = {
        "bienvenida.docx": {
            "title": "Bienvenida al Curso",
            "content": [
                "## Introducción",
                "¡Bienvenidos/as a este espacio de aprendizaje! En este curso exploraremos juntos los conceptos fundamentales de la materia, desarrollando competencias clave para su formación profesional.",
                "## Objetivos",
                "Al finalizar este curso, serás capaz de:",
                "- Comprender los conceptos teóricos fundamentales de la disciplina.",
                "- Aplicar metodologías de análisis en situaciones prácticas.",
                "## Metodología",
                "Este curso combina sesiones teóricas con actividades prácticas.",
            ],
        },
        "programa.docx": {
            "title": "Programa de la Asignatura",
            "content": [
                "## Información general",
                "Este programa detalla los contenidos, la carga horaria y los criterios de evaluación de la asignatura.",
                "## Contenidos mínimos",
                "- Fundamentos teóricos de la disciplina",
                "- Metodologías de análisis y aplicación",
                "- Herramientas y tecnologías contemporáneas",
                "- Proyectos de integración y casos de estudio",
                "## Bibliografía recomendada",
                "- Autor, A. (2024). Título del libro principal. Editorial.",
                "- Autor, B. (2023). Complemento teórico avanzado. Editorial.",
                "## Cronograma",
                "El curso se desarrolla en 16 semanas, divididas en los módulos detallados en la estructura del aula virtual.",
            ],
        },
        "tema1_intro.docx": {
            "title": "Introducción a los Conceptos Básicos",
            "content": [
                "En esta sección abordaremos los conceptos fundamentales que constituyen la base teórica de nuestra disciplina. Es esencial comprender estos principios antes de avanzar a temas más complejos.",
                "## ¿Qué son los conceptos básicos?",
                "Los conceptos básicos son los pilares sobre los que se construye todo el conocimiento posterior. Sin una comprensión sólida de estos fundamentos, resulta difícil avanzar en el aprendizaje.",
                "## Contexto histórico",
                "Estos conceptos fueron desarrollados a lo largo de décadas de investigación. Los principales aportes provienen de investigadores que sentaron las bases de la disciplina moderna.",
                "## Importancia en la práctica profesional",
                "El dominio de estos conceptos te permitirá:",
                "- Analizar problemas de forma estructurada",
                "- Proponer soluciones fundamentadas",
                "- Comunicar ideas con precisión técnica",
            ],
        },
        "tema1_desarrollo.docx": {
            "title": "Desarrollo Teórico",
            "content": [
                "## Marco teórico",
                "El marco teórico de esta unidad se sustenta en las principales corrientes de pensamiento contemporáneas. A continuación, exploraremos cada una de ellas en detalle.",
                "### Corriente principal",
                "La corriente principal establece que los principios fundamentales se basan en la observación sistemática y la experimentación controlada.",
                "### Perspectivas complementarias",
                "Existen perspectivas complementarias que enriquecen la comprensión del fenómeno, aportando visiones desde diferentes ángulos disciplinares.",
                "## Aplicaciones prácticas",
                "Los conceptos presentados tienen aplicaciones directas en diversos campos profesionales. A continuación, se detallan algunos ejemplos representativos.",
                "- Aplicación en el ámbito organizacional",
                "- Aplicación en la investigación académica",
                "- Aplicación en el desarrollo tecnológico",
            ],
        },
        "actividad1.docx": {
            "title": "Actividad 1: Ejercicio Práctico",
            "content": [
                "## Consigna",
                "Realiza el siguiente ejercicio práctico para aplicar los conceptos aprendidos en el Módulo 1.",
                "## Instrucciones",
                "1. Lee atentamente el material teórico del módulo.",
                "1. Identifica los conceptos clave presentados.",
                "1. Elabora un mapa conceptual que relacione al menos 5 conceptos.",
                "1. Redacta un párrafo de reflexión sobre la importancia de estos conceptos en tu campo profesional.",
                "## Criterios de evaluación",
                "- Identificación correcta de conceptos (25%)",
                "- Relaciones establecidas en el mapa conceptual (25%)",
                "- Coherencia y profundidad de la reflexión (30%)",
                "- Presentación y formato (20%)",
                "## Formato de entrega",
                "Sube tu trabajo en formato PDF a la sección de Tareas. Fecha límite: según calendario del curso.",
            ],
        },
        "tema2_analisis.docx": {
            "title": "Análisis Avanzado",
            "content": [
                "## Profundización en el análisis",
                "En este módulo profundizaremos en las técnicas de análisis avanzado, aplicando los conceptos fundamentales aprendidos en el módulo anterior.",
                "## Metodología de análisis",
                "La metodología propuesta se basa en un enfoque sistemático que permite abordar problemas complejos de forma estructurada.",
                "### Paso 1: Identificación del problema",
                "El primer paso consiste en definir claramente el problema a analizar, delimitando su alcance y estableciendo los objetivos del análisis.",
                "### Paso 2: Recolección de datos",
                "Una vez definido el problema, se procede a recolectar los datos necesarios utilizando las técnicas apropiadas.",
                "### Paso 3: Procesamiento e interpretación",
                "Los datos recolectados se procesan y analizan para extraer conclusiones fundamentadas.",
            ],
        },
        "tema2_casos.docx": {
            "title": "Casos de Estudio",
            "content": [
                "## Caso 1: Aplicación en contexto real",
                "Este caso presenta una situación real donde se aplicaron los conceptos y metodologías estudiadas, con resultados significativos para la organización involucrada.",
                "La organización enfrentaba el desafío de optimizar sus procesos internos. Mediante la aplicación de las técnicas de análisis estudiadas, se logró una mejora del 35% en la eficiencia operativa.",
                "## Caso 2: Investigación aplicada",
                "En este segundo caso se analiza un proyecto de investigación que utilizó las herramientas teóricas presentadas en el curso.",
                "Los investigadores diseñaron un estudio longitudinal que permitió validar los principios teóricos en un contexto específico, contribuyendo al avance del conocimiento en la disciplina.",
                "## Reflexión",
                "Estos casos demuestran la relevancia práctica de los contenidos abordados y la importancia de una formación sólida en los fundamentos de la disciplina.",
            ],
        },
        "actividad2.docx": {
            "title": "Actividad 2: Trabajo Integrador",
            "content": [
                "## Consigna",
                "Elabora un trabajo integrador que aplique los conceptos y metodologías de los Módulos 1 y 2.",
                "## Modalidad",
                "Este es un trabajo grupal (3-4 integrantes). Coordinen con sus compañeros/as a través del foro del módulo.",
                "## Instrucciones",
                "1. Seleccionen un problema real relacionado con su campo profesional.",
                "1. Apliquen la metodología de análisis estudiada.",
                "1. Desarrollen una propuesta de solución fundamentada.",
                "1. Presenten los resultados en un informe escrito.",
                "## Estructura del informe",
                "- Introducción y planteamiento del problema",
                "- Marco teórico (referencias a los contenidos del curso)",
                "- Metodología aplicada",
                "- Resultados y análisis",
                "- Conclusiones y recomendaciones",
                "- Bibliografía",
                "## Criterios de evaluación",
                "- Pertinencia del problema seleccionado (15%)",
                "- Aplicación correcta de la metodología (25%)",
                "- Calidad del análisis y fundamentación teórica (30%)",
                "- Propuesta de solución innovadora (20%)",
                "- Presentación y redacción (10%)",
            ],
        },
    }

    for filename, data in docx_files.items():
        create_example_docx(
            INPUT_DIR / filename,
            data["title"],
            data["content"],
        )

    print()
    print("=" * 50)
    print("  [OK] Archivos de ejemplo creados exitosamente")
    print()
    print("  Ahora ejecuta:")
    print('    python main.py --titulo "Curso de Ejemplo"')
    print("=" * 50)
    print()


if __name__ == "__main__":
    main()
