# -*- coding: utf-8 -*-
"""Análisis exploratorio de los casos reales: convenciones de carpetas,
formato de XLSX de estructura y estilos de los DOCX de módulos."""
import sys, io, re
from pathlib import Path
import openpyxl
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

CASOS = Path("casos")

def analizar_xlsx(p: Path):
    try:
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
        ws = wb.active
        filas = []
        for row in ws.iter_rows(values_only=True):
            filas.append([str(c).strip() if c is not None else "" for c in row])
        wb.close()
        modulos = [f[0] for f in filas if f and f[0] and "módulo" in f[0].lower()]
        secciones = [f[1] for f in filas if len(f) > 1 and f[1] and f[1].isupper() and len(f[1]) < 40]
        return {"filas": len(filas), "cols": max((len(f) for f in filas), default=0),
                "header": filas[0][:5] if filas else [], "modulos": len(modulos),
                "secciones": secciones[:10]}
    except Exception as e:
        return {"error": str(e)}

def analizar_docx(p: Path):
    try:
        d = Document(str(p))
        estilos = {}
        numerados = 0
        pat = re.compile(r"^\d+(\.\d+)+\.?\s")
        for para in d.paragraphs:
            t = para.text.strip()
            if not t:
                continue
            s = para.style.name if para.style else "?"
            if "Heading" in s or "Título" in s or "Title" in s:
                estilos[s] = estilos.get(s, 0) + 1
            if pat.match(t):
                bold = any(r.bold for r in para.runs if r.text.strip())
                if "Heading" in s or bold:
                    numerados += 1
        return {"headings": estilos, "titulos_numerados": numerados,
                "parrafos": len([x for x in d.paragraphs if x.text.strip()])}
    except Exception as e:
        return {"error": str(e)}

for curso in sorted(CASOS.iterdir()):
    if not curso.is_dir():
        continue
    print("=" * 90)
    print("CURSO:", curso.name)
    # Raíz real (los zips tienen una carpeta raíz adentro)
    root = curso
    hijos = list(root.iterdir())
    if len(hijos) == 1 and hijos[0].is_dir():
        root = hijos[0]
        hijos = list(root.iterdir())
        if len(hijos) == 1 and hijos[0].is_dir():
            root = hijos[0]

    print("  Carpetas de primer nivel:")
    for h in sorted(root.iterdir()):
        if h.is_dir():
            n = len(list(h.rglob("*")))
            print(f"    [{n:3d}] {h.name}")

    # XLSX de estructura
    estructuras = [p for p in root.rglob("*.xlsx") if "estructura" in p.name.lower()]
    for e in estructuras:
        info = analizar_xlsx(e)
        print(f"  XLSX estructura: {e.relative_to(root)}")
        print(f"    -> {info}")

    # DOCX de módulos (los archivos grandes de material multimedial / desarrollo)
    docx_mod = [p for p in root.rglob("*.docx")
                if re.search(r"(m[óo]dulo|multimedial|-M\d|_M\d| M\d)", p.name, re.I)
                and "borrador" not in str(p).lower() and "plantilla" not in p.name.lower()
                and "video" not in p.name.lower() and not p.name.startswith("~")]
    for m in sorted(docx_mod):
        info = analizar_docx(m)
        print(f"  DOCX módulo: {m.name[:70]}")
        print(f"    -> {info}")
