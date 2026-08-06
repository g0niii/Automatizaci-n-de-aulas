"""
Lector de archivos XLSX con instrucciones de estructura del curso.
Lee la hoja de cálculo que define módulos, páginas, orden y tipos de contenido.
"""

import logging
from pathlib import Path
from dataclasses import dataclass, field
from typing import Optional

import openpyxl

logger = logging.getLogger(__name__)


@dataclass
class CourseItem:
    """Representa un ítem individual dentro de un módulo del curso."""
    modulo: str
    orden: int
    tipo: str                             # pagina, archivo, url_externa, subencabezado
    titulo: str
    archivo_docx: Optional[str] = None
    archivo_media: list = field(default_factory=list)
    plantilla: str = "base_page"
    url: Optional[str] = None
    puntos: Optional[float] = 10.0
    intentos: Optional[int] = 1
    vencimiento: Optional[str] = None
    banner: Optional[str] = None
    notas: Optional[str] = None


@dataclass
class CourseModule:
    """Representa un módulo completo del curso con sus ítems."""
    nombre: str
    items: list = field(default_factory=list)  # Lista de CourseItem


def read_structure(filepath: Path) -> list[CourseModule]:
    """
    Lee el archivo XLSX de instrucciones y devuelve una lista de módulos
    con sus ítems ordenados.

    Args:
        filepath: Ruta al archivo .xlsx con la estructura del curso.

    Returns:
        Lista de CourseModule con sus ítems organizados.

    Raises:
        FileNotFoundError: Si el archivo no existe.
        ValueError: Si el formato del archivo es inválido.
    """
    if not filepath.exists():
        raise FileNotFoundError(f"No se encontró el archivo de estructura: {filepath}")

    logger.info(f"Leyendo estructura del curso desde: {filepath}")

    wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)
    ws = wb.active

    if ws is None:
        raise ValueError("El archivo XLSX no tiene una hoja activa.")

    # Leer encabezados de la primera fila
    headers = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        val = str(cell.value).strip().lower() if cell.value else ""
        headers.append(val)

    logger.debug(f"Encabezados encontrados: {headers}")

    # Mapear columnas esperadas
    required_cols = {"modulo", "orden", "tipo", "titulo"}
    found_cols = set(headers)
    missing = required_cols - found_cols
    if missing:
        raise ValueError(
            f"Faltan columnas requeridas en el XLSX: {missing}. "
            f"Columnas encontradas: {found_cols}"
        )

    # Índices de columnas
    col_idx = {name: idx for idx, name in enumerate(headers)}

    # Leer filas de datos
    modules_dict: dict[str, CourseModule] = {}
    row_count = 0

    for row in ws.iter_rows(min_row=2):
        values = [cell.value for cell in row]

        # Ignorar filas vacías
        modulo_val = values[col_idx["modulo"]]
        if modulo_val is None or str(modulo_val).strip() == "":
            continue

        modulo_name = str(modulo_val).strip()
        orden_val = values[col_idx["orden"]]
        tipo_val = str(values[col_idx["tipo"]]).strip().lower() if values[col_idx["tipo"]] else "pagina"
        titulo_val = str(values[col_idx["titulo"]]).strip() if values[col_idx["titulo"]] else ""

        # Campos opcionales
        archivo_docx = None
        if "archivo_docx" in col_idx and values[col_idx["archivo_docx"]]:
            archivo_docx = str(values[col_idx["archivo_docx"]]).strip()

        archivo_media = []
        if "archivo_media" in col_idx and values[col_idx["archivo_media"]]:
            media_str = str(values[col_idx["archivo_media"]]).strip()
            archivo_media = [m.strip() for m in media_str.split(";") if m.strip()]

        plantilla = "base_page"
        if "plantilla" in col_idx and values[col_idx["plantilla"]]:
            plantilla = str(values[col_idx["plantilla"]]).strip()

        url = None
        if "url" in col_idx and values[col_idx["url"]]:
            url = str(values[col_idx["url"]]).strip()

        notas = None
        if "notas" in col_idx and values[col_idx["notas"]]:
            notas = str(values[col_idx["notas"]]).strip()

        puntos = 10.0
        if "puntos" in col_idx and values[col_idx["puntos"]] is not None:
            try:
                puntos = float(values[col_idx["puntos"]])
            except ValueError:
                puntos = 10.0

        intentos = 1
        if "intentos" in col_idx and values[col_idx["intentos"]] is not None:
            try:
                intentos = int(values[col_idx["intentos"]])
            except ValueError:
                intentos = 1

        vencimiento = None
        if "vencimiento" in col_idx and values[col_idx["vencimiento"]]:
            vencimiento = str(values[col_idx["vencimiento"]]).strip()

        banner = None
        if "banner" in col_idx and values[col_idx["banner"]]:
            banner = str(values[col_idx["banner"]]).strip()

        item = CourseItem(
            modulo=modulo_name,
            orden=int(orden_val) if orden_val else row_count + 1,
            tipo=tipo_val,
            titulo=titulo_val,
            archivo_docx=archivo_docx,
            archivo_media=archivo_media,
            plantilla=plantilla,
            url=url,
            puntos=puntos,
            intentos=intentos,
            vencimiento=vencimiento,
            banner=banner,
            notas=notas,
        )

        # Agrupar por módulo
        if modulo_name not in modules_dict:
            modules_dict[modulo_name] = CourseModule(nombre=modulo_name)

        modules_dict[modulo_name].items.append(item)
        row_count += 1

    wb.close()

    # Ordenar ítems dentro de cada módulo
    for module in modules_dict.values():
        module.items.sort(key=lambda x: x.orden)

    modules = list(modules_dict.values())
    logger.info(f"Estructura leída: {len(modules)} módulos, {row_count} ítems totales.")

    return modules
